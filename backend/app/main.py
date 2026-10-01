from __future__ import annotations

import io
import os
import re
import uuid
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
load_dotenv()

from fastapi import Depends, FastAPI, Header, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from jinja2 import Environment, FileSystemLoader, select_autoescape
from pydantic import BaseModel, Field
from supabase import create_client, Client
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from weasyprint import HTML

from backend.services.resume_service import load_seed, resume_to_text
from backend.services.groq_service import (
    ats_original, ats_tailored, create_tailored_resume, cover_letter,
)

ROOT = Path(__file__).resolve().parents[2]
app = FastAPI(title='Resume Optimizer API', version='1.0.0')

cors = os.getenv('CORS_ORIGINS', '*')
app.add_middleware(CORSMiddleware, allow_origins=[x.strip() for x in cors.split(',')], allow_credentials=True, allow_methods=['*'], allow_headers=['*'])
app.mount('/assets', StaticFiles(directory=str(ROOT / 'frontend' / 'assets')), name='assets')
app.mount('/static', StaticFiles(directory=str(ROOT / 'frontend')), name='static')

jinja = Environment(loader=FileSystemLoader(str(ROOT / 'backend' / 'templates')), autoescape=select_autoescape(['html']))

SUPABASE_URL = os.getenv('SUPABASE_URL', '')
SUPABASE_ANON_KEY = os.getenv('SUPABASE_ANON_KEY', '')
SUPABASE_SERVICE_ROLE_KEY = os.getenv('SUPABASE_SERVICE_ROLE_KEY', '')
ADMIN_EMAILS = {x.strip().lower() for x in os.getenv('ADMIN_EMAILS', '').split(',') if x.strip()}


def db() -> Client:
    if not SUPABASE_URL or not SUPABASE_SERVICE_ROLE_KEY:
        raise HTTPException(500, 'Supabase server configuration is missing')
    return create_client(SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY)


def current_user(authorization: str | None = Header(default=None)) -> dict[str, Any]:
    if not authorization or not authorization.lower().startswith('bearer '):
        raise HTTPException(401, 'Authentication required')
    token = authorization.split(' ', 1)[1].strip()
    try:
        client = create_client(SUPABASE_URL, SUPABASE_ANON_KEY)
        result = client.auth.get_user(token)
        user = result.user
        if not user:
            raise ValueError('No authenticated user')
        email = (user.email or '').lower()
        if ADMIN_EMAILS and email not in ADMIN_EMAILS:
            raise HTTPException(403, 'This account is not authorized for Resume Optimizer')
        return {'id': user.id, 'email': user.email}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(401, f'Invalid or expired session: {exc}') from exc


def ensure_resume(user_id: str) -> dict[str, Any]:
    client = db()
    result = client.table('resume_configs').select('*').eq('user_id', user_id).limit(1).execute()
    if result.data:
        return result.data[0]
    resume = load_seed()
    inserted = client.table('resume_configs').insert({'user_id': user_id, 'resume': resume}).execute()
    return inserted.data[0]


class ResumePayload(BaseModel):
    resume: dict[str, Any]

class TailorPayload(BaseModel):
    job_description: str = Field(min_length=30)
    name: str = Field(default='Tailored Resume', min_length=1, max_length=120)

class CoverLetterPayload(BaseModel):
    job_description: str = Field(min_length=30)
    resume_id: str | None = None
    name: str = Field(default='Cover Letter', min_length=1, max_length=120)

class ATSOriginalPayload(BaseModel):
    resume: dict[str, Any] | None = None

class ATSTailoredPayload(BaseModel):
    resume_id: str
    job_description: str = Field(min_length=30)

class RenamePayload(BaseModel):
    name: str = Field(min_length=1, max_length=120)


def normalize_ats_report(report: dict[str, Any]) -> dict[str, Any]:
    # Keep LLM output aligned with the fixed rubric even if the model returns a
    # slightly inconsistent aggregate score. Component scores are clamped to their weights.
    total = 0.0
    for component in report.get('components', []):
        weight = float(component.get('weight', 0) or 0)
        score = min(weight, max(0.0, float(component.get('score', 0) or 0)))
        component['weight'] = weight
        component['score'] = round(score, 2)
        total += score
    report['score'] = round(min(100.0, max(0.0, total)), 2)
    return report


@app.get('/', response_class=HTMLResponse)
def index():
    return FileResponse(str(ROOT / 'frontend' / 'index.html'))

@app.get('/api/config.js')
def config_js():
    return HTMLResponse(content=f"window.__SUPABASE_URL__={SUPABASE_URL!r}; window.__SUPABASE_ANON_KEY__={SUPABASE_ANON_KEY!r};", media_type='application/javascript')

@app.get('/api/health')
def health():
    return {'ok': True, 'model': os.getenv('GROQ_MODEL', 'openai/gpt-oss-20b')}

@app.get('/api/me')
def me(user=Depends(current_user)):
    return user

@app.get('/api/resume/original')
def get_original(user=Depends(current_user)):
    return ensure_resume(user['id'])

@app.put('/api/resume/original')
def update_original(payload: ResumePayload, user=Depends(current_user)):
    client = db(); existing = ensure_resume(user['id'])
    version = int(existing.get('version', 1)) + 1
    result = client.table('resume_configs').update({'resume': payload.resume, 'version': version}).eq('user_id', user['id']).execute()
    return result.data[0]

@app.get('/api/resumes')
def list_resumes(user=Depends(current_user)):
    result = db().table('generated_resumes').select('id,name,job_description,ats_score,created_at,updated_at').eq('user_id', user['id']).order('created_at', desc=True).execute()
    return result.data

@app.get('/api/resumes/{resume_id}')
def get_resume(resume_id: str, user=Depends(current_user)):
    result = db().table('generated_resumes').select('*').eq('id', resume_id).eq('user_id', user['id']).single().execute()
    return result.data

@app.patch('/api/resumes/{resume_id}')
def rename_resume(resume_id: str, payload: RenamePayload, user=Depends(current_user)):
    result = db().table('generated_resumes').update({'name': payload.name}).eq('id', resume_id).eq('user_id', user['id']).execute()
    if not result.data: raise HTTPException(404, 'Resume not found')
    return result.data[0]

@app.delete('/api/resumes/{resume_id}')
def delete_resume(resume_id: str, user=Depends(current_user)):
    result = db().table('generated_resumes').delete().eq('id', resume_id).eq('user_id', user['id']).execute()
    return {'deleted': bool(result.data)}

@app.post('/api/resumes/tailor')
def tailor_resume(payload: TailorPayload, user=Depends(current_user)):
    original = ensure_resume(user['id'])['resume']
    tailored = create_tailored_resume(original, payload.job_description)
    client = db()
    inserted = client.table('generated_resumes').insert({
        'user_id': user['id'], 'name': payload.name, 'job_description': payload.job_description, 'resume': tailored
    }).execute()
    row = inserted.data[0]
    client.table('generation_events').insert({'user_id': user['id'], 'event_type':'resume_generated', 'job_title': payload.name, 'metadata': {'resume_id': row['id']}}).execute()
    return row

@app.post('/api/ats/original')
def score_original(payload: ATSOriginalPayload, user=Depends(current_user)):
    resume = payload.resume or ensure_resume(user['id'])['resume']
    report = normalize_ats_report(ats_original(resume))
    db().table('ats_reports').insert({'user_id': user['id'], 'report_type':'original', 'score':report['score'], 'report':report}).execute()
    db().table('generation_events').insert({'user_id': user['id'], 'event_type':'ats_original', 'metadata': {'score': report['score']}}).execute()
    return report

@app.post('/api/ats/tailored')
def score_tailored(payload: ATSTailoredPayload, user=Depends(current_user)):
    row = db().table('generated_resumes').select('*').eq('id', payload.resume_id).eq('user_id', user['id']).single().execute().data
    if not row: raise HTTPException(404, 'Resume not found')
    report = normalize_ats_report(ats_tailored(row['resume'], payload.job_description))
    client = db()
    client.table('ats_reports').insert({'user_id': user['id'], 'report_type':'tailored', 'resume_id':row['id'], 'job_description':payload.job_description, 'score':report['score'], 'report':report}).execute()
    client.table('generated_resumes').update({'ats_score':report['score']}).eq('id',row['id']).execute()
    client.table('generation_events').insert({'user_id': user['id'], 'event_type':'ats_tailored', 'metadata': {'score': report['score'], 'resume_id':row['id']}}).execute()
    return report

@app.post('/api/cover-letters')
def make_cover_letter(payload: CoverLetterPayload, user=Depends(current_user)):
    client = db()
    resume = None
    if payload.resume_id:
        found = client.table('generated_resumes').select('*').eq('id', payload.resume_id).eq('user_id', user['id']).single().execute()
        resume = found.data
    if not resume:
        # If no generated resume is selected, create from the current canonical resume for a safe baseline.
        resume = ensure_resume(user['id'])
    content = cover_letter(resume['resume'], payload.job_description)
    inserted = client.table('cover_letters').insert({'user_id':user['id'], 'name':payload.name, 'job_description':payload.job_description, 'resume_id':payload.resume_id, 'content':content}).execute()
    client.table('generation_events').insert({'user_id':user['id'], 'event_type':'cover_letter_generated', 'metadata': {'cover_letter_id':inserted.data[0]['id']}}).execute()
    return inserted.data[0]

@app.get('/api/cover-letters')
def list_cover_letters(user=Depends(current_user)):
    return db().table('cover_letters').select('id,name,job_description,resume_id,created_at,updated_at').eq('user_id',user['id']).order('created_at',desc=True).execute().data

@app.get('/api/cover-letters/{letter_id}')
def get_cover_letter(letter_id: str, user=Depends(current_user)):
    return db().table('cover_letters').select('*').eq('id',letter_id).eq('user_id',user['id']).single().execute().data

@app.patch('/api/cover-letters/{letter_id}')
def rename_cover_letter(letter_id: str, payload: RenamePayload, user=Depends(current_user)):
    result=db().table('cover_letters').update({'name':payload.name}).eq('id',letter_id).eq('user_id',user['id']).execute()
    if not result.data: raise HTTPException(404,'Cover letter not found')
    return result.data[0]

@app.get('/api/analytics')
def analytics(user=Depends(current_user)):
    client=db()
    resumes=client.table('generated_resumes').select('id,name,job_description,ats_score,created_at').eq('user_id',user['id']).order('created_at',desc=True).execute().data
    letters=client.table('cover_letters').select('id,created_at').eq('user_id',user['id']).execute().data
    ats=client.table('ats_reports').select('score,report_type,created_at').eq('user_id',user['id']).order('created_at',desc=True).execute().data
    events=client.table('generation_events').select('event_type,job_title,company,metadata,created_at').eq('user_id',user['id']).order('created_at',desc=True).limit(100).execute().data
    scores=[float(x['score']) for x in ats if x.get('score') is not None]
    return {
      'generated_count':len(resumes), 'cover_letter_count':len(letters), 'best_ats_score':max(scores) if scores else None,
      'ats_runs':len(ats), 'recent_resumes':resumes[:8], 'recent_events':events[:12],
      'score_history':ats[:12],
    }


def render_resume(resume: dict[str, Any]) -> str:
    return jinja.get_template('resume.html').render(resume=resume)


def build_docx(resume: dict[str, Any]) -> bytes:
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
    styles=doc.styles; styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(9)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(resume['personal_information']['name']); r.bold=True; r.font.size=Pt(20)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(resume['personal_information']['title']); r.font.size=Pt(10.5)
    c=resume['personal_information'].get('contact',{}); links=[]
    if c.get('email'): links.append(c['email'])
    if c.get('website'): links.append('Portfolio')
    if c.get('github'): links.append('GitHub')
    if c.get('linkedin'): links.append('LinkedIn')
    if c.get('x'): links.append('X')
    if c.get('hugging_face'): links.append('Hugging Face')
    if c.get('dev_to'): links.append('Dev.to')
    p=doc.add_paragraph(' · '.join(links)); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    def heading(t):
        p=doc.add_paragraph(); r=p.add_run(t.upper()); r.bold=True; r.font.size=Pt(10); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(2)
    def bullets(items):
        for item in items:
            p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(1); p.add_run(item)
    s=resume.get('professional_summary',{}).get('summary','')
    if s: heading('Professional Summary'); doc.add_paragraph(s)
    heading('Skills')
    for k,v in resume.get('skills',{}).items():
        if v: p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); r=p.add_run(k.replace('_',' ').title()+': '); r.bold=True; p.add_run(', '.join(v))
    if resume.get('experience'):
        heading('Experience')
        for x in resume['experience']:
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); r=p.add_run(f"{x.get('position','')} · {x.get('company','')}"); r.bold=True; p.add_run(f"  |  {x.get('start_date','')} – {x.get('end_date','')}")
            bullets(x.get('responsibilities',[])+[f"Achievement: {a}" for a in x.get('achievements',[])])
    if resume.get('projects'):
        heading('Projects')
        for x in resume['projects']:
            p=doc.add_paragraph(); r=p.add_run(x.get('name','')); r.bold=True
            bullets(x.get('description',[]))
            if x.get('technical_skills'): doc.add_paragraph('Technologies: '+', '.join(x['technical_skills']))
            if x.get('github') or x.get('live_url'): doc.add_paragraph('Links: '+' · '.join([u for u in [x.get('github'),x.get('live_url')] if u]))
    for key,title in [('education','Education'),('certifications','Certifications'),('achievements','Achievements'),('languages','Languages')]:
        items=resume.get(key,[])
        if items:
            heading(title)
            for x in items:
                if key=='education': text=f"{x.get('degree','')} — {x.get('field','')} · {x.get('institution','')} · {x.get('start_date','')} – {x.get('end_date','')} · {x.get('grade',{}).get('value','')}"
                elif key=='certifications': text=f"{x.get('name','')} · {x.get('issuer','')} · {x.get('date','')}"
                elif key=='achievements': text=f"{x.get('title','')} · {x.get('organization','')} · {x.get('date','')}"
                else: text=f"{x.get('language','')} · {x.get('level','')} · {x.get('proficiency','')}"
                doc.add_paragraph(text.strip(' ·'))
    out=io.BytesIO(); doc.save(out); return out.getvalue()


def filename_safe(name: str, ext: str) -> str:
    safe=re.sub(r'[^a-zA-Z0-9._-]+','_',name).strip('_') or 'resume'
    return safe[:100]+ext

@app.get('/api/resumes/{resume_id}/download/{fmt}')
def download_resume(resume_id: str, fmt: str, user=Depends(current_user)):
    row=db().table('generated_resumes').select('*').eq('id',resume_id).eq('user_id',user['id']).single().execute().data
    return download_resume_data(row['name'],row['resume'],fmt)

@app.get('/api/resume/original/download/{fmt}')
def download_original(fmt: str,user=Depends(current_user)):
    row=ensure_resume(user['id']); return download_resume_data('Current_Resume',row['resume'],fmt)

def download_resume_data(name: str,resume:dict[str,Any],fmt:str):
    fmt=fmt.lower()
    if fmt=='pdf':
        pdf=HTML(string=render_resume(resume),base_url=str(ROOT)).write_pdf()
        return StreamingResponse(io.BytesIO(pdf),media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="{filename_safe(name,".pdf")}"'})
    if fmt=='docx':
        data=build_docx(resume); return StreamingResponse(io.BytesIO(data),media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':f'attachment; filename="{filename_safe(name,".docx")}"'})
    raise HTTPException(400,'Format must be pdf or docx')

@app.get('/api/cover-letters/{letter_id}/download/{fmt}')
def download_letter(letter_id:str,fmt:str,user=Depends(current_user)):
    row=db().table('cover_letters').select('*').eq('id',letter_id).eq('user_id',user['id']).single().execute().data
    fmt=fmt.lower()
    if fmt=='pdf':
        html=f'<html><head><style>@page{{size:A4;margin:22mm}}body{{font-family:Arial;font-size:11pt;line-height:1.55;color:#202938}}h1{{font-size:20pt}}p{{white-space:pre-wrap}}</style></head><body><h1>{row["name"]}</h1><p>{row["content"]}</p></body></html>'
        data=HTML(string=html).write_pdf(); return StreamingResponse(io.BytesIO(data),media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="{filename_safe(row["name"],".pdf")}"'})
    if fmt=='docx':
        doc=Document(); doc.styles['Normal'].font.name='Arial'; doc.styles['Normal'].font.size=Pt(11); doc.add_heading(row['name'],0); doc.add_paragraph(row['content']); out=io.BytesIO(); doc.save(out)
        return StreamingResponse(io.BytesIO(out.getvalue()),media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':f'attachment; filename="{filename_safe(row["name"],".docx")}"'})
    raise HTTPException(400,'Format must be pdf or docx')
